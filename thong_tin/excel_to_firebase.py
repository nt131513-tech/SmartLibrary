import json
import sys
from pathlib import Path
from openpyxl import load_workbook


def val(x):
    """Chuyển ô rỗng thành chuỗi rỗng."""
    return "" if x is None else x


def rows(ws):
    """Đọc dữ liệu từ sheet, dòng đầu tiên là tên cột."""
    headers = [val(c.value) for c in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(x not in (None, "") for x in row):
            yield {
                str(headers[i]).strip(): val(row[i])
                for i in range(len(headers))
                if headers[i] not in (None, "")
            }


def text(x):
    """Chuyển dữ liệu sang chuỗi, tránh xuất hiện chuỗi 'None'."""
    return "" if x is None else str(x).strip()


def integer(x):
    try:
        return int(x)
    except (TypeError, ValueError):
        return 0


def check_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Khong tim thay sheet '{sheet_name}'. "
            f"Cac sheet hien co: {', '.join(wb.sheetnames)}"
        )


def convert_books(wb):
    check_sheet(wb, "Books")
    books = {}

    for x in rows(wb["Books"]):
        book_id = text(x.get("id"))

        if not book_id:
            continue

        quantity = integer(x.get("quantity"))
        total = integer(x.get("totalQuantity")) or quantity

        books[book_id] = {
            "title": text(x.get("title")),
            "author": text(x.get("author")),
            "category": text(x.get("category")),
            "isbn": text(x.get("isbn")),
            "shelf": text(x.get("shelf")),
            "quantity": quantity,
            "totalQuantity": total,
            "available": quantity > 0,
            "barcode": text(x.get("barcode")),
            "description": text(x.get("description")),
        }

    return books


def convert_accounts(wb, sheet_name, expected_role):
    """
    Chuyển sheet Người_dùng hoặc Quản_thư thành:
    - accounts: dictionary theo UID
    - accounts_without_uid: danh sách chưa có UID
    """
    check_sheet(wb, sheet_name)

    accounts = {}
    accounts_without_uid = []

    for x in rows(wb[sheet_name]):
        record = {
            key: value
            for key, value in x.items()
            if key not in ("uid", "role") and value not in (None, "")
        }

        role = text(x.get("role")).lower() or expected_role

        if role != expected_role:
            raise ValueError(
                f"Role khong hop le trong sheet '{sheet_name}': "
                f"{text(x.get('email'))}. "
                f"Role phai la '{expected_role}'."
            )

        record["role"] = expected_role
        uid = text(x.get("uid"))

        if uid:
            accounts[uid] = record
        else:
            accounts_without_uid.append(record)

    return accounts, accounts_without_uid


def convert_borrow_records(wb):
    check_sheet(wb, "BorrowRecords")
    borrow_records = {}

    for x in rows(wb["BorrowRecords"]):
        record_id = text(x.get("id"))

        if not record_id:
            continue

        borrow_records[record_id] = {
            "userId": text(x.get("userId")),
            "bookId": text(x.get("bookId")),
            "borrowDate": text(x.get("borrowDate")),
            "returnDate": text(x.get("returnDate")),
            "status": text(x.get("status")),
        }

    return borrow_records


def load_json_if_exists(path):
    """Đọc JSON nếu tồn tại; nếu không có thì trả về dict rỗng."""
    if not path.exists():
        return {}

    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise ValueError(f"File JSON khong hop le: {path}") from error

    if not isinstance(data, dict):
        raise ValueError(f"File {path} phai co dang object JSON.")

    return data


def merge_accounts(existing_accounts, new_accounts):
    """
    Gộp dữ liệu cũ và mới theo UID.

    - UID cũ không xuất hiện trong Excel vẫn được giữ nguyên.
    - UID có trong Excel sẽ được cập nhật bằng dữ liệu mới.
    - Các trường cũ không có trong dữ liệu mới vẫn được giữ lại.
    """
    merged = dict(existing_accounts)

    for uid, new_record in new_accounts.items():
        old_record = merged.get(uid, {})

        if isinstance(old_record, dict):
            merged[uid] = {**old_record, **new_record}
        else:
            merged[uid] = new_record

    return merged


def save_json(output_folder, filename, data):
    output_path = output_folder / filename
    output_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    print(f"Da tao: {output_path} ({len(data)} dong)")


def main():
    if len(sys.argv) != 2:
        print("Cach dung:")
        print("py excel_to_firebase_separate_safe.py SmartLibrary_Data.xlsx")
        raise SystemExit(1)

    excel_path = Path(sys.argv[1])

    if not excel_path.exists():
        raise FileNotFoundError(f"Khong tim thay file Excel: {excel_path}")

    wb = load_workbook(excel_path, data_only=True)

    output_folder = Path("firebase_import")
    output_folder.mkdir(exist_ok=True)

    # Chuyển dữ liệu từ Excel
    books = convert_books(wb)

    users_from_excel, users_without_uid = convert_accounts(
        wb, "Người_dùng", "user"
    )

    librarians_from_excel, librarians_without_uid = convert_accounts(
        wb, "Quản_thư", "librarian"
    )

    borrow_records = convert_borrow_records(wb)

    # Người dùng và quản thư dùng chung node /users trong Firebase.
    accounts_from_excel = {
        **users_from_excel,
        **librarians_from_excel,
    }

    # Ưu tiên file export trực tiếp từ Firebase:
    # firebase_import/users_existing.json
    existing_users_path = output_folder / "users_existing.json"

    # Nếu chưa có file export, dùng users.json cũ làm dữ liệu nền.
    # Tuy nhiên, users_existing.json vẫn được khuyến nghị.
    if existing_users_path.exists():
        existing_users = load_json_if_exists(existing_users_path)
        print(f"Da doc du lieu cu tu: {existing_users_path}")
    else:
        old_users_path = output_folder / "users.json"
        existing_users = load_json_if_exists(old_users_path)

        if existing_users:
            print(
                "Canh bao: Dang dung users.json cu lam du lieu nen. "
                "Nen export node /users tren Firebase thanh "
                "users_existing.json truoc khi import."
            )

    # Merge: giữ tài khoản cũ, thêm tài khoản mới, cập nhật UID trùng.
    merged_users = merge_accounts(existing_users, accounts_from_excel)

    # Xuất các file riêng
    save_json(output_folder, "books.json", books)
    save_json(output_folder, "users.json", merged_users)
    save_json(output_folder, "librarians.json", librarians_from_excel)
    save_json(output_folder, "borrowRecords.json", borrow_records)
    save_json(output_folder, "users_without_uid.json", users_without_uid)
    save_json(
        output_folder,
        "librarians_without_uid.json",
        librarians_without_uid,
    )

    print("\nHoan tat!")
    print("users.json da gom ca user va librarian theo UID.")
    print("Khi import vao Firebase, hay import users.json vao node /users.")

    if users_without_uid:
        print(
            f"Can bo sung UID cho {len(users_without_uid)} nguoi dung "
            "trong Firebase Authentication."
        )

    if librarians_without_uid:
        print(
            f"Can bo sung UID cho {len(librarians_without_uid)} quan thu "
            "trong Firebase Authentication."
        )

    print("\nQUAN TRONG:")
    print("1. Truoc khi cap nhat, export node /users tren Firebase.")
    print("2. Luu file export voi ten: firebase_import/users_existing.json")
    print("3. Chay lai script de tao users.json da merge.")
    print("4. Import users.json vao /users; khong import truc tiep du lieu Excel chua merge.")


if __name__ == "__main__":
    main()
